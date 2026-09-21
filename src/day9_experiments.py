"""Checkpointed Day 9 experiment orchestration.

The runner reuses the audited Day 5--8 pricing engines.  Every replication is
appended immediately, so an interrupted long run can resume without changing
the frozen seed schedule.
"""

from __future__ import annotations

import gc
import json
import math
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, MutableMapping

import numpy as np
import openpyxl
import pandas as pd

from autocallable_bb import BridgeBank
from autocallable_direct import DirectContract, equicorrelation, make_psd_correlation, parse_research_contract, stable_seed
from autocallable_gpu import (
    conditioned_replication_shared_input,
    configure_cupy_cache,
    direct_replication_shared_input,
    generate_endpoint_normal_input,
    generate_normal_input,
    gpu_environment,
    warmup_gpu,
)
from autocallable_greeks import autocall_redemption_replication
from day9_analysis import sha256_file


RESULT_FIELDS = (
    "total_value",
    "fair_coupon",
    "coupon_value",
    "early_redemption_principal",
    "surviving_notional",
    "terminal_ki_loss",
    "continuous_ki_survival_probability",
    "component_identity_error",
    "probability_mass_error",
    "fair_coupon_residual",
)


@dataclass(frozen=True)
class MarketInputs:
    tickers: tuple[str, ...]
    initial_spot_ratios: np.ndarray
    dividend_yields: np.ndarray
    volatilities: np.ndarray
    correlation: np.ndarray
    risk_free_rate: float
    workbook_as_of: pd.Timestamp
    rate_as_of: pd.Timestamp
    workbook_sha256: str


@dataclass(frozen=True)
class Scenario:
    scenario_id: str
    regime_family: str
    regime_level: str
    contract: DirectContract
    risk_free_rate: float
    dividend_yields: np.ndarray
    volatilities: np.ndarray
    correlation: np.ndarray
    initial_spot_ratios: np.ndarray


def _is_number(value: object) -> bool:
    return (
        isinstance(value, (int, float, np.integer, np.floating))
        and not isinstance(value, bool)
        and np.isfinite(value)
    )


def load_market_inputs(project_root: str | Path, core_config: dict[str, Any]) -> MarketInputs:
    root = Path(project_root)
    workbook_path = root / core_config["market_data"]["relative_path"]
    workbook_hash = sha256_file(workbook_path)
    expected_hash = core_config["market_data"]["sha256"]
    if workbook_hash != expected_hash:
        raise AssertionError("frozen workbook hash mismatch")
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=False)
    setup = workbook["Setup"]
    snapshot = workbook["Underlying_Snapshot"]
    history = workbook["Underlying_History"]
    market_history = workbook["Market_History"]
    issuer_initial = workbook["Issuer_Initial_Value"]
    tickers = tuple(str(snapshot.cell(row, 2).value) for row in (6, 7, 8))
    spots = np.array([snapshot.cell(row, 4).value for row in (6, 7, 8)], dtype=float)
    dividend_yields = np.array(
        [snapshot.cell(row, 5).value for row in (6, 7, 8)], dtype=float
    ) / 100.0
    volatilities = np.array(
        [snapshot.cell(row, 8).value for row in (6, 7, 8)], dtype=float
    ) / 100.0
    initial_references = np.array(
        [issuer_initial.cell(row, 3).value for row in (38, 39, 40)], dtype=float
    )
    histories = []
    for date_column, value_column, ticker in zip((1, 4, 7), (2, 5, 8), tickers):
        values: dict[pd.Timestamp, float] = {}
        for row in range(6, history.max_row + 1):
            date_value = history.cell(row, date_column).value
            level = history.cell(row, value_column).value
            if hasattr(date_value, "year") and _is_number(level) and float(level) > 0:
                values[pd.Timestamp(date_value)] = float(level)
        histories.append(pd.Series(values, name=ticker).sort_index())
    prices = pd.concat(histories, axis=1, join="inner").dropna()
    correlation = make_psd_correlation(
        np.log(prices / prices.shift(1)).dropna().corr().to_numpy()
    )
    rate_rows = []
    for row in range(6, market_history.max_row + 1):
        date_value = market_history.cell(row, 10).value
        rate_pct = market_history.cell(row, 11).value
        if hasattr(date_value, "year") and _is_number(rate_pct) and float(rate_pct) > 0:
            rate_rows.append((pd.Timestamp(date_value), float(rate_pct) / 100.0))
    rates = pd.Series(dict(rate_rows)).sort_index()
    return MarketInputs(
        tickers=tickers,
        initial_spot_ratios=spots / initial_references,
        dividend_yields=dividend_yields,
        volatilities=volatilities,
        correlation=correlation,
        risk_free_rate=float(rates.iloc[-1]),
        workbook_as_of=pd.Timestamp(setup["B8"].value),
        rate_as_of=pd.Timestamp(rates.index[-1]),
        workbook_sha256=workbook_hash,
    )


def _append_checkpoint(path: Path, row: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame([row]).to_csv(
        path,
        mode="a",
        header=not path.exists(),
        index=False,
        float_format="%.17g",
    )


def _existing_keys(path: Path, columns: Iterable[str]) -> set[tuple[Any, ...]]:
    if not path.exists() or path.stat().st_size == 0:
        return set()
    frame = pd.read_csv(path)
    return set(frame.loc[:, list(columns)].itertuples(index=False, name=None))


def _method_replication(
    scenario: Scenario,
    method: str,
    n_paths: int,
    replication: int,
    base_seed: int,
    direct_steps_per_year: int,
    direct_batch_size: int,
    inner_paths: int,
    bridge_substeps: int,
    bessel_terms: int,
    bridge_bank_seed: int,
    bank_cache: MutableMapping[tuple[Any, ...], BridgeBank],
) -> dict[str, Any]:
    seed = stable_seed(base_seed, scenario.scenario_id, method, n_paths, replication)
    if method in {"M0", "M1"}:
        random_method = "mc" if method == "M0" else "rqmc"
        grid_steps = int(math.ceil(scenario.contract.maturity_years * direct_steps_per_year))
        # monitoring_grid may add exact observation nodes; the engine validates the shape.
        from autocallable_direct import monitoring_grid

        grid, _ = monitoring_grid(scenario.contract, direct_steps_per_year)
        normal_input = generate_normal_input(n_paths, len(grid) - 1, random_method, seed)
        result, _ = direct_replication_shared_input(
            contract=scenario.contract,
            risk_free_rate=scenario.risk_free_rate,
            dividend_yields=scenario.dividend_yields,
            volatilities=scenario.volatilities,
            correlation=scenario.correlation,
            annual_coupon=scenario.contract.baseline_annual_coupon,
            n_paths=n_paths,
            steps_per_year=direct_steps_per_year,
            method=random_method,
            seed=seed,
            backend="gpu",
            batch_size=min(direct_batch_size, n_paths),
            initial_spot_ratios=scenario.initial_spot_ratios,
            normal_input=normal_input,
        )
        random_input_seconds = float(normal_input.generation_seconds)
        runtime_total_seconds = float(result["total_with_input_seconds"])
        del normal_input
    elif method in {"M2", "M3"}:
        outer_method = "mc" if method == "M2" else "rqmc"
        endpoint_input = generate_endpoint_normal_input(
            n_paths, len(scenario.contract.observation_times), outer_method, seed
        )
        result = conditioned_replication_shared_input(
            contract=scenario.contract,
            risk_free_rate=scenario.risk_free_rate,
            dividend_yields=scenario.dividend_yields,
            volatilities=scenario.volatilities,
            correlation=scenario.correlation,
            annual_coupon=scenario.contract.baseline_annual_coupon,
            endpoint_normals=endpoint_input.values,
            outer_method=outer_method,
            seed=seed,
            bridge_bank_seed=bridge_bank_seed,
            inner_paths=inner_paths,
            bridge_substeps=bridge_substeps,
            bessel_terms=bessel_terms,
            initial_spot_ratios=scenario.initial_spot_ratios,
            endpoint_backend="gpu",
            probability_backend="gpu-hybrid",
            bank_cache=bank_cache,
            return_diagnostics=False,
        )
        random_input_seconds = float(endpoint_input.generation_seconds)
        runtime_total_seconds = float(result["total_wall_seconds"] + random_input_seconds)
        del endpoint_input
    else:
        raise KeyError(method)

    row = {
        "scenario_id": scenario.scenario_id,
        "regime_family": scenario.regime_family,
        "regime_level": scenario.regime_level,
        "contract_id": scenario.contract.contract_id,
        "method": method,
        "n_paths": int(n_paths),
        "replication": int(replication),
        "seed": int(seed),
        "random_input_seconds": random_input_seconds,
        "runtime_engine_seconds": float(result["total_wall_seconds"]),
        "runtime_total_seconds": runtime_total_seconds,
        "risk_free_rate": float(scenario.risk_free_rate),
        "volatility_1": float(scenario.volatilities[0]),
        "volatility_2": float(scenario.volatilities[1]),
        "volatility_3": float(scenario.volatilities[2]),
        "correlation_12": float(scenario.correlation[0, 1]),
        "correlation_13": float(scenario.correlation[0, 2]),
        "correlation_23": float(scenario.correlation[1, 2]),
        "ki_barrier_ratio": float(scenario.contract.ki_barrier_ratio),
        "initial_spot_ratio_1": float(scenario.initial_spot_ratios[0]),
        "initial_spot_ratio_2": float(scenario.initial_spot_ratios[1]),
        "initial_spot_ratio_3": float(scenario.initial_spot_ratios[2]),
        "direct_steps_per_year": int(direct_steps_per_year) if method in {"M0", "M1"} else np.nan,
        "inner_paths": int(inner_paths) if method in {"M2", "M3"} else np.nan,
        "bridge_substeps": int(bridge_substeps) if method in {"M2", "M3"} else np.nan,
        "bessel_terms": int(bessel_terms) if method in {"M2", "M3"} else np.nan,
        "backend": "gpu" if method in {"M0", "M1"} else "gpu-hybrid",
    }
    for field in RESULT_FIELDS:
        row[field] = float(result[field])
    gc.collect()
    return row


def _baseline_scenario(contract: DirectContract, market: MarketInputs) -> Scenario:
    return Scenario(
        scenario_id="baseline_rc_a",
        regime_family="baseline",
        regime_level="asymmetric_hsbc_informed",
        contract=contract,
        risk_free_rate=market.risk_free_rate,
        dividend_yields=market.dividend_yields,
        volatilities=market.volatilities,
        correlation=market.correlation,
        initial_spot_ratios=market.initial_spot_ratios,
    )


def build_regime_scenarios(
    core_config: dict[str, Any],
    day9_config: dict[str, Any],
    market: MarketInputs,
) -> list[Scenario]:
    rc_a = parse_research_contract(core_config, "RC-A")
    rc_l = parse_research_contract(core_config, "RC-L")
    scenarios: list[Scenario] = []
    correlation_levels = {
        "low": equicorrelation(float(day9_config["regimes"]["correlation"]["low_equicorrelation"])),
        "baseline": market.correlation,
        "high": equicorrelation(float(day9_config["regimes"]["correlation"]["high_equicorrelation"])),
    }
    for level, correlation in correlation_levels.items():
        scenarios.append(
            Scenario(
                f"correlation_{level}",
                "correlation",
                level,
                rc_a,
                market.risk_free_rate,
                market.dividend_yields,
                market.volatilities,
                correlation,
                market.initial_spot_ratios,
            )
        )
    for scale in day9_config["regimes"]["volatility_scales"]:
        label = f"{int(round(float(scale) * 100))}%"
        scenarios.append(
            Scenario(
                f"volatility_{label}",
                "volatility",
                label,
                rc_a,
                market.risk_free_rate,
                market.dividend_yields,
                market.volatilities * float(scale),
                market.correlation,
                market.initial_spot_ratios,
            )
        )
    for level, barrier in day9_config["regimes"]["ki_barriers"].items():
        contract = rc_a.with_overrides(
            contract_id=f"RC-A-KI-{str(level).upper()}",
            label=f"{rc_a.label} - {level} KI regime",
            ki_barrier_ratio=float(barrier),
        )
        scenarios.append(
            Scenario(
                f"ki_barrier_{level}",
                "continuous_KI_barrier",
                str(level),
                contract,
                market.risk_free_rate,
                market.dividend_yields,
                market.volatilities,
                market.correlation,
                market.initial_spot_ratios,
            )
        )
    scenarios.extend(
        [
            Scenario(
                "parameters_symmetric_lee",
                "parameterisation",
                "symmetric_lee",
                rc_l,
                0.03,
                np.zeros(3),
                np.full(3, 0.20),
                equicorrelation(0.40),
                np.ones(3),
            ),
            Scenario(
                "parameters_asymmetric_hsbc_informed",
                "parameterisation",
                "asymmetric_hsbc_informed",
                rc_a,
                market.risk_free_rate,
                market.dividend_yields,
                market.volatilities,
                market.correlation,
                market.initial_spot_ratios,
            ),
        ]
    )
    return scenarios


def warmup_day9(
    scenario: Scenario,
    config: dict[str, Any],
) -> dict[str, Any]:
    started = time.perf_counter()
    environment = warmup_gpu()
    baseline = config["baseline"]
    cache: dict[tuple[Any, ...], BridgeBank] = {}
    _method_replication(
        scenario,
        "M0",
        512,
        -1,
        config["seed"],
        baseline["direct_steps_per_year"],
        baseline["direct_batch_size"],
        baseline["conditioned_inner_paths"],
        baseline["conditioned_bridge_substeps"],
        baseline["bessel_terms"],
        stable_seed(config["seed"], "warmup-bank"),
        cache,
    )
    _method_replication(
        scenario,
        "M3",
        256,
        -1,
        config["seed"],
        baseline["direct_steps_per_year"],
        baseline["direct_batch_size"],
        baseline["conditioned_inner_paths"],
        baseline["conditioned_bridge_substeps"],
        baseline["bessel_terms"],
        stable_seed(config["seed"], "warmup-bank"),
        cache,
    )
    return {
        **gpu_environment(),
        "compile_and_pipeline_warmup_seconds": time.perf_counter() - started,
        "warmup_status": environment.get("warmup_status", "PASS"),
    }


def run_baseline_and_reference(
    project_root: str | Path,
    core_config: dict[str, Any],
    config: dict[str, Any],
    market: MarketInputs,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    root = Path(project_root)
    output_dir = root / config["evidence_directory"]
    baseline_path = output_dir / "baseline_replications.csv"
    reference_path = output_dir / "reference_replications.csv"
    scenario = _baseline_scenario(parse_research_contract(core_config, config["contract_id"]), market)
    baseline_cfg = config["baseline"]
    reference_cfg = config["reference"]
    baseline_keys = _existing_keys(baseline_path, ("method", "n_paths", "replication"))
    bank_cache: dict[tuple[Any, ...], BridgeBank] = {}
    completed = len(baseline_keys)
    total = len(baseline_cfg["methods"]) * len(baseline_cfg["path_grid"]) * baseline_cfg["replications"]
    for n_paths in baseline_cfg["path_grid"]:
        for method in baseline_cfg["methods"]:
            for replication in range(baseline_cfg["replications"]):
                key = (method, n_paths, replication)
                if key in baseline_keys:
                    continue
                row = _method_replication(
                    scenario,
                    method,
                    n_paths,
                    replication,
                    config["seed"],
                    baseline_cfg["direct_steps_per_year"],
                    baseline_cfg["direct_batch_size"],
                    baseline_cfg["conditioned_inner_paths"],
                    baseline_cfg["conditioned_bridge_substeps"],
                    baseline_cfg["bessel_terms"],
                    stable_seed(config["seed"], "baseline-bridge-bank"),
                    bank_cache,
                )
                _append_checkpoint(baseline_path, row)
                baseline_keys.add(key)
                completed += 1
                if completed % 8 == 0 or completed == total:
                    print(f"baseline checkpoint {completed}/{total}", flush=True)

    reference_keys = _existing_keys(reference_path, ("method", "n_paths", "replication"))
    reference_cache: dict[tuple[Any, ...], BridgeBank] = {}
    reference_total = reference_cfg["replications"]
    for replication in range(reference_cfg["replications"]):
        key = (reference_cfg["method"], reference_cfg["n_paths"], replication)
        if key in reference_keys:
            continue
        row = _method_replication(
            scenario,
            reference_cfg["method"],
            reference_cfg["n_paths"],
            replication,
            stable_seed(config["seed"], "reference"),
            baseline_cfg["direct_steps_per_year"],
            baseline_cfg["direct_batch_size"],
            reference_cfg["conditioned_inner_paths"],
            reference_cfg["conditioned_bridge_substeps"],
            reference_cfg["bessel_terms"],
            stable_seed(config["seed"], "reference-bridge-bank"),
            reference_cache,
        )
        _append_checkpoint(reference_path, row)
        reference_keys.add(key)
        print(f"reference checkpoint {len(reference_keys)}/{reference_total}", flush=True)
    return pd.read_csv(baseline_path), pd.read_csv(reference_path)


def run_regimes(
    project_root: str | Path,
    core_config: dict[str, Any],
    config: dict[str, Any],
    market: MarketInputs,
) -> pd.DataFrame:
    root = Path(project_root)
    output_path = root / config["evidence_directory"] / "regime_replications.csv"
    regime_cfg = config["regimes"]
    keys = _existing_keys(output_path, ("scenario_id", "method", "replication"))
    scenarios = build_regime_scenarios(core_config, config, market)
    baseline_cfg = config["baseline"]
    total = len(scenarios) * len(regime_cfg["methods"]) * regime_cfg["replications"]
    completed = len(keys)
    for scenario in scenarios:
        bank_cache: dict[tuple[Any, ...], BridgeBank] = {}
        for method in regime_cfg["methods"]:
            for replication in range(regime_cfg["replications"]):
                key = (scenario.scenario_id, method, replication)
                if key in keys:
                    continue
                row = _method_replication(
                    scenario,
                    method,
                    regime_cfg["n_paths"],
                    replication,
                    stable_seed(config["seed"], "regimes"),
                    baseline_cfg["direct_steps_per_year"],
                    baseline_cfg["direct_batch_size"],
                    baseline_cfg["conditioned_inner_paths"],
                    baseline_cfg["conditioned_bridge_substeps"],
                    baseline_cfg["bessel_terms"],
                    stable_seed(config["seed"], scenario.scenario_id, "bridge-bank"),
                    bank_cache,
                )
                _append_checkpoint(output_path, row)
                keys.add(key)
                completed += 1
                if completed % 16 == 0 or completed == total:
                    print(f"regime checkpoint {completed}/{total}", flush=True)
    return pd.read_csv(output_path)


def run_trigger_smoothing(
    project_root: str | Path,
    core_config: dict[str, Any],
    config: dict[str, Any],
    market: MarketInputs,
) -> pd.DataFrame:
    root = Path(project_root)
    output_path = root / config["evidence_directory"] / "trigger_smoothing_replications.csv"
    trigger_cfg = config["trigger_smoothing"]
    keys = _existing_keys(
        output_path,
        ("worst_spot_ratio", "time_to_observation_days", "smooth", "replication"),
    )
    contract = parse_research_contract(core_config, config["contract_id"])
    total = (
        len(trigger_cfg["worst_spot_grid"])
        * len(trigger_cfg["time_to_observation_days"])
        * 2
        * trigger_cfg["replications"]
    )
    completed = len(keys)
    for days in trigger_cfg["time_to_observation_days"]:
        for worst_spot in trigger_cfg["worst_spot_grid"]:
            spots = np.array(
                [
                    trigger_cfg["other_spot_ratios"][0],
                    trigger_cfg["other_spot_ratios"][1],
                    worst_spot,
                ],
                dtype=float,
            )
            for smooth in (False, True):
                for replication in range(trigger_cfg["replications"]):
                    key = (worst_spot, days, smooth, replication)
                    if key in keys:
                        continue
                    seed = stable_seed(
                        config["seed"], "trigger", worst_spot, days, replication
                    )
                    result = autocall_redemption_replication(
                        initial_spot_ratios=spots,
                        trigger_ratios=trigger_cfg["trigger_ratio"],
                        time_to_observation=float(days) / 365.0,
                        principal=contract.principal,
                        risk_free_rate=market.risk_free_rate,
                        dividend_yields=market.dividend_yields,
                        volatilities=market.volatilities,
                        correlation=market.correlation,
                        n_paths=trigger_cfg["n_paths"],
                        method=trigger_cfg["method"],
                        seed=seed,
                        smooth=smooth,
                    )
                    row = {
                        **result,
                        "worst_spot_ratio": float(worst_spot),
                        "time_to_observation_days": int(days),
                        "replication": int(replication),
                        "scope": trigger_cfg["scope"],
                    }
                    _append_checkpoint(output_path, row)
                    keys.add(key)
                    completed += 1
                    if completed % 32 == 0 or completed == total:
                        print(f"trigger checkpoint {completed}/{total}", flush=True)
    return pd.read_csv(output_path)


def run_all_experiments(
    project_root: str | Path,
    core_config: dict[str, Any],
    config: dict[str, Any],
) -> dict[str, Any]:
    root = Path(project_root)
    output_dir = root / config["evidence_directory"]
    output_dir.mkdir(parents=True, exist_ok=True)
    configure_cupy_cache(output_dir / ".cupy_cache")
    market = load_market_inputs(root, core_config)
    scenario = _baseline_scenario(parse_research_contract(core_config, config["contract_id"]), market)
    warmup_record = warmup_day9(scenario, config)
    pd.DataFrame([warmup_record]).to_csv(output_dir / "environment.csv", index=False)
    baseline, reference = run_baseline_and_reference(root, core_config, config, market)
    regimes = run_regimes(root, core_config, config, market)
    trigger = run_trigger_smoothing(root, core_config, config, market)
    market_frame = pd.DataFrame(
        {
            "ticker": market.tickers,
            "initial_spot_ratio": market.initial_spot_ratios,
            "dividend_yield": market.dividend_yields,
            "volatility": market.volatilities,
        }
    )
    market_frame.to_csv(output_dir / "market_inputs.csv", index=False)
    pd.DataFrame(market.correlation, index=market.tickers, columns=market.tickers).to_csv(
        output_dir / "correlation_matrix.csv"
    )
    return {
        "market": market,
        "warmup": warmup_record,
        "baseline_replications": baseline,
        "reference_replications": reference,
        "regime_replications": regimes,
        "trigger_replications": trigger,
    }
